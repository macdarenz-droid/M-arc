"""Builds the upload files for M/ARC's Wear Engine application 3.

The Wear Engine form accepts only .xlsx uploads (it refused PDFs on 2026-09-29), so the uploads are
Huawei's own templates, filled in:
  MARC_Data_Permission_and_Usage.xlsx   from "Data Permission and Usage Description for Individual
                                        Developers.xlsx" (linked from Huawei's "Applying for the
                                        Wear Engine Service" page)
  MARC_User_Authorization_Path.xlsx     from "User Authorization Path Description.xlsx" (the form's
                                        Download Example)
The pictures inside them are rendered from two PDFs this script also builds; page 3 of each PDF
(the C2 watch screens) is copied unchanged from the application-2 PDFs.

The wording here is the source of truth; docs/WATCH-WEAR-ENGINE-APPLICATION.md section 5 summarises
it. Nothing personal is stored in this file. The xlsx files contain no name or email. --name and
--email only add a developer row to page 1 of the data-permission PDF, which is not uploaded.

Needs: python3 with reportlab, pypdf, pymupdf, openpyxl and Pillow, and the DejaVu Sans fonts.

  python3 docs/wear-engine/make_uploads.py --date "29 September 2026" \
    --old-dpu MARC_Data_Permission_and_Usage.pdf --old-aup MARC_User_Authorization_Path.pdf \
    --dpu-template individual.xlsx --aup-template auth.xlsx \
    --train train.jpg --live live.jpg --watch watch.jpg --out OUT_DIR
"""
import argparse
import io
import os
import sys
from xml.sax.saxutils import escape

from pypdf import PdfReader, PdfWriter
from reportlab.lib.utils import ImageReader

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.dont_write_bytecode = True
from pdfkit import (A4, ACCENT, CELL, INK, L, MUTED, R, RULE, TEXTW, H, Page, Paragraph, canvas,
                    style)

SHA = '05:66:9A:D2:72:1C:6A:BA:F9:FD:D4:B9:B8:4E:2F:B7:94:48:44:B1:DE:F3:59:84:5F:01:5F:2B:67:CA:F1:F5'
APP_ROWS = [
    ('Phone app', 'M/ARC (Android)'),
    ('App ID and package', '119100049<br/>com.mrcdrnzz.dailytracker'),
    ('Wearable target', 'HUAWEI WATCH GT 6<br/>Harmony Lightweight Smart Wearable Device'),
]
# Huawei's own wording for this permission, from its individual-developer template.
BASIC_INFO = ("Obtain paired wearable devices' random identifiers, names, battery levels, connection "
              "statuses, app installation statuses, and other information, as well as send audio and "
              "other files to the devices.")
FOOT_DPU = 'M/ARC  |  App ID 119100049  |  Permission and usage'
FOOT_AUP = 'M/ARC  |  App ID 119100049  |  User authorization path'


def header(p, developer=None):
    rows = list(APP_ROWS)
    if developer:
        rows.append(developer)
    p.kv(rows)
    p.gap(17)
    p.para('SHA-256 signing certificate', 9.2, bold=True)
    p.gap(4)
    p.para(SHA, 9.0)


def image(c, path, x, y_top, w, h=None, border=True):
    img = ImageReader(path)
    iw, ih = img.getSize()
    h = h or w * ih / iw
    c.drawImage(img, x, y_top - h, w, h)
    if border:
        c.setStrokeColor(RULE); c.setLineWidth(0.8); c.rect(x, y_top - h, w, h, stroke=1, fill=0)
    return h


def caption(c, html, x, y_top, w, size=8.6, color=INK):
    p = Paragraph(html, style(size, color, leading=11.6))
    _, h = p.wrap(w, 400)
    p.drawOn(c, x, y_top - h)
    return h


def disclosure_mock(c, x, y_top, w, h):
    """Proposed M/ARC connect disclosure, drawn in the app's light theme."""
    c.setFillColor(CELL); c.setStrokeColor(RULE); c.setLineWidth(1.0)
    c.roundRect(x, y_top - h, w, h, 14, stroke=1, fill=0)
    c.setFillColorRGB(0.965, 0.957, 0.937)
    c.roundRect(x + 5, y_top - h + 5, w - 10, h - 10, 11, stroke=0, fill=1)
    inner = x + 14
    iw = w - 28
    y = y_top - 22
    c.setFillColor(MUTED); c.setFont('DV', 7.2); c.drawString(inner, y, 'SETTINGS  ·  WATCH')
    y -= 16
    y -= caption(c, '<b>Connect HUAWEI WATCH GT 6</b>', inner, y, iw, 10.5)
    y -= 8
    body = ('M/ARC sends your current exercise, set and rest timer to the M/ARC watch app, and '
            'receives the sets and rest actions you confirm on the watch.<br/><br/>'
            'Your workout records are stored only on this phone. Keep M/ARC open during the workout.'
            '<br/><br/>Uses Huawei Health and HUAWEI Wear Engine.')
    y -= caption(c, body, inner, y, iw, 7.6, MUTED)
    by = y_top - h + 52
    c.setFillColorRGB(0.176, 0.165, 0.149)
    c.roundRect(inner, by, iw, 22, 9, stroke=0, fill=1)
    c.setFillColorRGB(1, 1, 1); c.setFont('DVB', 8.4); c.drawCentredString(inner + iw / 2, by + 7.5, 'Continue')
    c.setFillColor(INK); c.setFont('DV', 8.4); c.drawCentredString(inner + iw / 2, by - 18, 'Not now')


def huawei_screen_card(c, x, y_top, w, h):
    c.setFillColor(CELL); c.setStrokeColor(RULE); c.setLineWidth(1.0)
    c.setDash(4, 3)
    c.roundRect(x, y_top - h, w, h, 14, stroke=1, fill=1)
    c.setDash()
    inner = x + 12
    y = y_top - 20
    y -= caption(c, '<b>Huawei Health authorization screen</b>', inner, y, w - 24, 9.6)
    y -= 8
    body = ('Shown by Huawei Health, not by M/ARC, when M/ARC requests Basic device information '
            'and it is not granted yet.<br/><br/>'
            'With Huawei Health 11.0.3.512 or later, Basic device information is granted by default, '
            "so this screen may not appear (Huawei guide: Requesting User Authorization).<br/><br/>"
            'It cannot be captured before this application is approved: until then the Wear Engine '
            'API returns error code 8.')
    caption(c, body, inner, y, w - 24, 7.9, MUTED)


# ---------------------------------------------------------------- Data Permission and Usage
def build_dpu(a, path):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)

    p = Page(c, 'Data Permission and Usage Description',
             f'Wear Engine application | Individual developer | {a.date}', FOOT_DPU, 1, 4)
    header(p, ('Individual developer', f'{escape(a.name)}<br/>{escape(a.email)}') if a.name else None)
    p.heading("Self-check (Huawei's individual-developer template)", 20, 8)
    p.kv([
        ('1  Phone app type', 'Android app, as selected on the request page. Non-Huawei Android phones are '
                              'supported where Huawei Health runs in the background. There is no iOS app.'),
        ('2  Develop a wearable app?', 'Yes. The M/ARC watch app for HUAWEI WATCH GT 6 (a lite wearable app) '
                                       'is in development; page 3 shows its screens.'),
        ('3  Watch app talks to the phone app?', 'Yes, through Wear Engine P2P messages (page 2).'),
        ('4  Phone app stays active?', 'Yes. The M/ARC phone app UI stays in the foreground and active, '
                                       'Huawei Health runs in the background, and the M/ARC watch app is open '
                                       'in the foreground at the same time. If either app is not active, '
                                       'nothing is sent; the watch holds confirmed actions until M/ARC is open '
                                       'again (page 2).'),
        ('5  Wearable models', 'HUAWEI WATCH GT 6.'),
    ], label_w=150)
    p.heading('Permission requested', 18, 8)
    p.kv([('1  Basic device information',
           f'{BASIC_INFO}<br/><b>This is the only permission requested.</b> M/ARC uses it for the '
           'paired-device list and connection status (to select the GT 6 and show Connected or Not '
           'connected), the installation status of the M/ARC watch app (to prompt installation), and '
           'P2P messages between the M/ARC phone app and the M/ARC watch app. Battery level, audio and '
           'file transfer are not used. Device information stays on the phone.<br/>'
           '<b>Usage scenario:</b> page 2. <b>Data display path:</b> page 2 (data flow), page 3 (watch '
           'screens) and page 4 (real phone screenshots).')], label_w=150)
    p.done()

    p = Page(c, 'Usage scenario and data flow', 'Product design for Wear Engine P2P communication',
             FOOT_DPU, 2, 4)
    p.para('<b>Usage scenario:</b> when the user starts a workout in M/ARC with the phone app in the '
           'foreground, M/ARC checks that the paired GT 6 is connected and that the M/ARC watch app is '
           'installed, and asks the user to install it if it is not. The phone sends the current exercise, '
           'set, target and rest timer to the watch app. The user confirms each set and rest action on the '
           'watch; the watch sends it to the phone, which saves it in the workout and replies Saved.', 9.8)
    p.gap(6)
    p.para('<b>Requirements:</b> phone apps must run Android. Third-party Android phones need to be '
           'supported and iOS phones do not need to be supported. A lite wearable app is developed for the '
           'HUAWEI WATCH GT 6. The phone app runs in the foreground, and the watch app is opened at the same '
           'time and runs in the foreground.', 9.8)
    p.gap(16)
    p.flow_diagram(('M/ARC phone', 'Stores the workout record'), ('M/ARC watch app', 'Holds unsent actions'),
                   [('Workout snapshot', 'right'), ('Confirmed action with ID', 'left'),
                    ('Saved reply', 'right')])
    p.heading('Messages exchanged', 2, 6)
    p.para('<b>Phone to watch:</b> current exercise, target, set count, rest timer and theme.<br/>'
           '<b>Watch to phone:</b> complete set (weight, reps, effort), add 30 seconds of rest, skip rest, '
           'pause, resume, finish. Each action carries an ID and is sent only when the user confirms it. '
           'Taps and swipes are handled on the watch; no touch coordinates or gesture stream are sent.', 9.8)
    p.heading('When the phone app is not active', 14, 6)
    p.para('If M/ARC is not active mid-workout, the watch keeps the last workout state, runs the rest '
           'countdown locally and shows <b>Phone not connected</b>. Confirmed actions wait in a small '
           'on-watch queue marked <b>Pending</b>, never Saved. When M/ARC is open again the watch resends '
           'them and refreshes; the phone applies each action once by its ID and replies <b>Saved</b>. '
           'With no workout running, the watch shows <b>Start your workout in M/ARC on your phone</b>.', 9.8)
    p.heading('Why a phone connection, not the internet', 14, 6)
    p.para('M/ARC has no accounts and no workout server. Workout records are stored only on the '
           "user's phone, and logging works offline in gyms. The watch needs the live workout held by "
           'that phone. The only online feature is an optional AI coach that answers the user\'s '
           'questions; its relay stores no workout data and the watch link never uses it.', 9.8)
    p.gap(12)
    p.para('<b>Implementation status:</b> the phone-side handling that applies each action once by its ID '
           'is built and tested. The watch app, the foreground messaging and the Saved reply are designed '
           'and still need end-to-end testing on the GT 6.', 9.0, MUTED)
    p.done()

    # page 3 is copied from application 2; page 4 is drawn on a placeholder page
    p = Page(c, 'App information', 'Real screenshots of the current M/ARC phone app', FOOT_DPU, 4, 4)
    col = (TEXTW - 2 * 14) / 3
    top = p.y
    shots = [(a.train, '<b>Workouts:</b> the split and its exercises.'),
             (a.live, '<b>Live workout:</b> sets, weight, reps and effort are logged here.'),
             (a.watch, '<b>Watch today:</b> heart rate from the watch\'s standard heart-rate broadcast '
                       'over direct Bluetooth. This does not use Wear Engine.')]
    low = top
    for i, (img, cap) in enumerate(shots):
        x = L + i * (col + 14)
        h = image(c, img, x, top, col)
        ch = caption(c, cap, x, top - h - 6, col)
        low = min(low, top - h - 6 - ch)
    p.y = low - 16
    p.para('<b>App introduction:</b> M/ARC is an Android gym workout tracker. It plans training splits, '
           'logs exercises, sets, weight, reps and effort, times sessions and rest, and keeps workout '
           'history on the phone. These screenshots show the phone app as it is today; they do not show '
           'Wear Engine, which is not yet approved.', 9.4)
    p.gap(6)
    p.para('<b>AppGallery:</b> M/ARC is not yet listed on AppGallery, so it has no rating. It is in '
           "private development and testing on the developer's own Android phone and HUAWEI WATCH GT 6.", 9.4)
    p.gap(6)
    p.para('<b>Supplementary information:</b> 1. No purchase of Huawei hardware through Huawei is '
           'planned. 2. No co-marketing with Huawei is planned. 3. Industry: sports.', 9.4)
    p.done()
    c.save()
    merge(buf, a.old_dpu, path)


# ---------------------------------------------------------------- User Authorization Path
def build_aup(a, path):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)

    p = Page(c, 'User Authorization Path Description',
             f'Wear Engine application | Phone and watch flow | {a.date}', FOOT_AUP, 1, 4)
    header(p)
    p.gap(22)
    p.box('<b>Permission requested:</b> Basic device information only, for P2P messages between the '
          'M/ARC phone app and the M/ARC watch app. Only the phone app applies for Wear Engine.')
    p.heading('User authorization path', 20, 8)
    p.kv([
        ('1  Open the watch screen', 'In M/ARC the user opens Settings → Watch and health → Watch, or taps '
                                     'the heart-rate button on the workout screen, and chooses Connect '
                                     'HUAWEI WATCH GT 6 (a proposed option on this existing screen).'),
        ('2  M/ARC disclosure', 'M/ARC explains what the watch link exchanges and that workout records are '
                                'stored only on the phone. Continue proceeds; Not now keeps normal phone '
                                'logging (page 2, screen 2).'),
        ('3  Huawei authorization', 'M/ARC checks the Basic device information permission through Wear '
                                    'Engine. With Huawei Health 11.0.3.512 or later it is granted by default. '
                                    'Otherwise Huawei Health shows its own authorization screen, where the '
                                    'user allows it. If the user cancels, M/ARC stays in phone-only logging.'),
        ('4  Select the watch', 'M/ARC lists the paired devices, the user selects the GT 6, and M/ARC '
                                'checks that the M/ARC watch app is installed. If it is not, M/ARC asks the '
                                'user to install it on the watch.'),
        ('5  Use during a workout', 'The user starts the workout in M/ARC and opens the M/ARC watch app '
                                    "manually from the watch's app list. Both apps stay open; Huawei Health "
                                    'runs in the background. Watch actions show Pending until the phone '
                                    'replies Saved (page 4).'),
        ('6  Change or stop', 'Disconnect on the same Watch screen stops all watch messages. Ending a '
                              'workout needs a deliberate Finish. No background or always-on operation is '
                              'claimed.'),
    ], label_w=150)
    p.gap(14)
    p.para('<b>Heart rate:</b> today M/ARC reads the watch\'s standard heart-rate broadcast over direct '
           'Bluetooth, which does not use Wear Engine. No Wear Engine sensor permission is requested.',
           9.1, MUTED)
    p.done()

    p = Page(c, 'Authorization screens', 'Where the user grants and changes access', FOOT_AUP, 2, 4)
    col = (TEXTW - 2 * 14) / 3
    top = p.y
    h1 = image(c, a.watch, L, top, col)
    caption(c, '<b>1 · Watch screen in M/ARC today (real).</b> The GT 6 option is added here. '
               'Disconnect stops the link.', L, top - h1 - 6, col)
    mh = 300
    disclosure_mock(c, L + col + 14, top, col, mh)
    caption(c, '<b>2 · M/ARC disclosure</b> (proposed design, shown before any watch message).',
            L + col + 14, top - mh - 6, col)
    huawei_screen_card(c, L + 2 * (col + 14), top, col, mh)
    caption(c, "<b>3 · Huawei Health authorization</b> (Huawei's own screen).",
            L + 2 * (col + 14), top - mh - 6, col)
    p.y = top - mh - 40
    p.heading('Granting access', 10, 6)
    p.para('Screen 1 → Connect HUAWEI WATCH GT 6 → screen 2 → Continue → screen 3 if Huawei Health asks '
           '→ select the GT 6. Cancel at screen 2 or 3 leaves M/ARC in phone-only logging, and nothing is '
           'sent to the watch.', 9.8)
    p.heading('Changing or removing access', 14, 6)
    p.para('Screen 1 → Disconnect. M/ARC then sends and accepts no watch messages until the user connects '
           'again. Workout logging on the phone continues.', 9.8)
    p.done()

    # page 3 is copied from application 2
    p = Page(c, 'Phone inactive and reconnection', 'Proposed design | No delivery promise while the phone '
             'app is inactive', FOOT_AUP, 4, 4)
    steps = [
        ('Action confirmed on the watch', 'The watch stores the action in its small local queue and shows '
                                          'Pending.'),
        ('M/ARC is not active', 'The watch keeps the last workout state, runs the rest countdown locally and '
                                'shows Phone not connected. Actions stay Pending, never Saved.'),
        ('M/ARC is open in the foreground again', 'The watch resends automatically and refreshes from the '
                                                   'phone. The phone checks each action ID and applies it '
                                                   'once.'),
        ('The phone saves the action', 'The phone replies Saved. Only then does the watch replace Pending '
                                       'with Saved.'),
    ]
    for i, (t, b) in enumerate(steps):
        p.step_box(t, b, 70)
        if i < len(steps) - 1:
            p.down_arrow()
    p.gap(20)
    p.box('<b>No workout running</b><br/>Start your workout in M/ARC on your phone')
    p.heading('Local data', 18, 6)
    p.para("Workout records are stored only on the user's phone. M/ARC has no accounts or workout server. "
           'The on-watch queue holds only actions waiting for the phone\'s reply. The only online feature is '
           "an optional AI coach that answers the user's questions; its relay stores no workout data and the "
           'watch link never uses it. The watch app does not use the internet.', 9.8)
    p.done()
    c.save()
    merge(buf, a.old_aup, path)


def merge(buf, old_pdf, path):
    """Insert page 3 of the application-2 file (the C2 watch screens) as page 3."""
    new = PdfReader(io.BytesIO(buf.getvalue()))
    old = PdfReader(old_pdf)
    w = PdfWriter()
    w.add_page(new.pages[0]); w.add_page(new.pages[1]); w.add_page(old.pages[2]); w.add_page(new.pages[2])
    w.add_metadata({'/Title': os.path.splitext(os.path.basename(path))[0].replace('_', ' '),
                    '/Author': 'M/ARC', '/Creator': 'docs/wear-engine/make_uploads.py'})
    with open(path, 'wb') as f:
        w.write(f)


# ---------------------------------------------------------------- xlsx uploads
SELF_CHECK = [  # rows 3-7 of the template, column C ("Self-Check Result (M)")
    'Android app, the same as selected on the request page. Non-Huawei Android phones are supported '
    'where Huawei Health runs in the background. There is no iOS or HarmonyOS phone app.',
    'Yes. The M/ARC watch app for HUAWEI WATCH GT 6 (a lite wearable JS app) is in development. Its '
    'screens are shown as UX diagrams in the images under this table (see D9).',
    'Yes. The watch app and the M/ARC phone app exchange Wear Engine P2P messages: a small workout '
    'snapshot from phone to watch, and the workout actions the user confirms from watch to phone.',
    'Yes. The phone app UI stays in the foreground and active during the workout, Huawei Health runs in '
    'the background, and the M/ARC watch app is open in the foreground at the same time. If either app '
    'is not active, nothing is sent: the watch keeps working locally, shows Phone not connected and '
    'holds confirmed actions until M/ARC is open again.',
    'HUAWEI WATCH GT 6',
]
USAGE = (
    'Usage scenario: when the user starts a workout in M/ARC with the phone app in the foreground, M/ARC '
    'checks that the paired HUAWEI WATCH GT 6 is connected and that the M/ARC watch app is installed, and '
    'asks the user to install it if it is not. The phone sends the current exercise, set, target weight '
    'and reps, rest timer and theme to the watch app. The user confirms each action on the watch '
    '(complete set with weight, reps and effort; add 30 s of rest; skip rest; pause; resume; finish). The '
    'watch sends it to the phone as one small P2P message with an ID; the phone applies it once, saves it '
    'in the workout and replies Saved, and the watch shows Pending until then. Taps and swipes are handled '
    'on the watch; no touch coordinates or gesture stream are sent.\n'
    'Requirement: the phone app UI is in the foreground and stays active, and the watch app is opened by '
    "the user from the watch's app list, so both apps are active in the foreground. Huawei Health runs in "
    'the background.\n'
    'Used: the paired-device list and connection status, the M/ARC watch app installation status, and P2P '
    'messages. Not used: battery level, audio and file transfer. Device information stays on the phone.')
DISPLAY = ('Phone- and wearable-side images are placed under this table, from row 12:\n'
           '1. Watch app UX diagrams (HUAWEI WATCH GT 6, proposed design, sample data), including the '
           'Phone not connected and Saved states.\n'
           '2. Usage scenario and data flow between the phone app and the watch app.\n'
           '3. What happens when the phone app is not active, and after it reopens.\n'
           '4. Real screenshots of the current M/ARC phone app.')
REMARKS = ('Basic device information is the only permission requested.\n'
           'Note 3: a watch app is developed, so Basic device information alone covers the messages.\n'
           'Note 5 (why phone sync, not the cloud): M/ARC has no accounts and no workout server. Workout '
           "records are stored only on the user's phone, and logging works offline in gyms, so the watch "
           'needs the live workout held by that phone. The only online feature is an optional AI coach that '
           "answers the user's questions; its relay stores no workout data and the watch link never uses it.\n"
           'Status: the phone-side handling that applies each action once by its ID is built and tested. The '
           'watch app and its messaging still need end-to-end testing on the GT 6.')
APP_INFO = [
    'App introduction: M/ARC is an Android gym workout tracker by an individual developer. It plans training '
    'splits, logs exercises, sets, weight, reps and effort, times sessions and rest, and keeps workout '
    "history on the phone. Today it shows live heart rate from the watch's standard heart-rate broadcast "
    'over direct Bluetooth, which does not use Wear Engine.',
    "AppGallery rating: M/ARC is not yet listed on AppGallery, so it has no rating. It is in private "
    "development and testing on the developer's own Android phone and HUAWEI WATCH GT 6.",
    'Screenshots of the main function screens (real, current app): Workouts, Live workout, Watch connection.',
]
SUPPLEMENTARY = ["No. No Huawei hardware purchase through Huawei is planned; testing uses the developer's "
                 'own HUAWEI WATCH GT 6.',
                 'No co-marketing with Huawei is planned.',
                 'Sports']
AUTH_TEXT = (
    'Permission requested: Basic device information only, for P2P messages between the M/ARC phone app '
    'and the M/ARC watch app. For Android apps it is granted by default (this template, and Huawei Health '
    '11.0.3.512 or later).\n\n'
    'How the user reaches it (images in column C):\n'
    '1. In M/ARC the user opens Settings > Watch and health > Watch, or taps the heart-rate button on the '
    'workout screen, and chooses Connect HUAWEI WATCH GT 6 (a proposed option on this existing screen; '
    'screen 1).\n'
    '2. M/ARC explains what the watch link exchanges and that workout records are stored only on the '
    'phone. Continue proceeds; Not now keeps normal phone logging (screen 2, proposed design).\n'
    '3. M/ARC checks the Basic device information permission through Wear Engine. It is granted by default; '
    'if Huawei Health still asks, it shows its own authorization screen (screen 3). Cancel leaves M/ARC in '
    'phone-only logging and nothing is sent to the watch. This screen cannot be captured before approval, '
    'because the API returns error code 8 until then.\n'
    '4. M/ARC lists the paired devices, the user selects the GT 6, and M/ARC checks that the M/ARC watch '
    'app is installed; if it is not, M/ARC asks the user to install it on the watch.\n'
    "5. During a workout the user opens the M/ARC watch app from the watch's app list. Both apps stay open; "
    'Huawei Health runs in the background.')
AUTH_MODIFY = ('Changing or removing access: in M/ARC, Settings > Watch and health > Watch > Disconnect '
               '(screen 1). M/ARC then sends and accepts no watch messages until the user connects again; '
               'workout logging on the phone continues.')


def render(pdf, page, png, clip=None, zoom=1.6):
    import pymupdf
    d = pymupdf.open(pdf)
    rect = pymupdf.Rect(*clip) if clip else None
    d[page].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=rect).save(png)
    return png


def strip(paths, png, width=386, gap=24):
    """Screenshots side by side on white, same width, top-aligned."""
    from PIL import Image
    ims = [Image.open(x).convert('RGB') for x in paths]
    ims = [im.resize((width, round(im.height * width / im.width))) for im in ims]
    out = Image.new('RGB', (width * len(ims) + gap * (len(ims) - 1), max(im.height for im in ims)), 'white')
    for i, im in enumerate(ims):
        out.paste(im, (i * (width + gap), 0))
    out.save(png)
    return png


def put(ws, cell, text, height=None, vertical='top'):
    """Writes wrapped text in black; Huawei's template cells carry blue example styling."""
    from copy import copy
    from openpyxl.styles import Alignment
    ws[cell] = text
    ws[cell].alignment = Alignment(wrap_text=True, vertical=vertical)
    font = copy(ws[cell].font)
    font.color = '000000'
    ws[cell].font = font
    if height:
        ws.row_dimensions[ws[cell].row].height = height


def picture(ws, png, cell, width_px):
    from openpyxl.drawing.image import Image as XImage
    img = XImage(png)
    img.height = round(img.height * width_px / img.width)
    img.width = width_px
    ws.add_image(img, cell)


def build_xlsx(a, tmp):
    from openpyxl import load_workbook
    dpu_pdf = os.path.join(tmp, 'MARC_Data_Permission_and_Usage.pdf')
    aup_pdf = os.path.join(tmp, 'MARC_User_Authorization_Path.pdf')
    img = lambda n: os.path.join(tmp, n)
    ux = render(dpu_pdf, 2, img('1_watch_ux.png'), (30, 30, 565, 800))
    flow = render(dpu_pdf, 1, img('2_scenario_and_flow.png'), (30, 30, 565, 800))
    inactive = render(aup_pdf, 3, img('3_phone_inactive.png'), (30, 30, 565, 800))
    phone = strip([a.train, a.live, a.watch], img('4_phone_today.png'))
    screens = render(aup_pdf, 1, img('auth_screens.png'), (36, 100, 560, 458), 2.0)

    wb = load_workbook(a.dpu_template)
    ws = wb['Individual Developer (M)']
    for row, text in zip(range(3, 8), SELF_CHECK):
        put(ws, f'C{row}', text, max(ws.row_dimensions[row].height or 15, 16 * (len(text) // 60 + 1) + 6))
    ws.unmerge_cells('E9:E10')
    ws.delete_rows(10)  # Message notification: "Delete rows for permissions not needed."
    put(ws, 'C9', USAGE, 330)
    put(ws, 'D9', DISPLAY)
    put(ws, 'E9', REMARKS)
    put(ws, 'B11', 'Images for D9 (Data Display Path): 1 watch app UX, 2 usage scenario and data flow, '
                   '3 phone app not active, 4 current phone app.')
    for n, (png, col) in enumerate([(ux, 'B'), (flow, 'C'), (inactive, 'D')], 1):
        put(ws, f'{col}12', f'{n}.')
        picture(ws, png, f'{col}13', 440)
    put(ws, 'B47', '4.')  # below the three page images (about 32 rows tall)
    picture(ws, phone, 'B48', 640)
    wi = wb['App Information (M)']
    for row, text in enumerate(APP_INFO, 2):
        put(wi, f'A{row}', text, 48)
    picture(wi, phone, 'A6', 660)
    wsu = wb['Supplementary Info (M)']
    for row, text in enumerate(SUPPLEMENTARY, 2):
        put(wsu, f'B{row}', text, vertical='center')
    wb.properties.lastModifiedBy = 'M/ARC'
    wb.save(os.path.join(a.out, 'MARC_Data_Permission_and_Usage.xlsx'))

    wb = load_workbook(a.aup_template)
    ws = wb.active
    ws._images = []  # Huawei's example screenshot
    put(ws, 'B2', AUTH_TEXT, 330)
    picture(ws, screens, 'C2', 560)
    put(ws, 'A3', 'User authorization modification (optional)', vertical='center')
    put(ws, 'B3', AUTH_MODIFY, 60)
    wb.properties.lastModifiedBy = 'M/ARC'
    wb.save(os.path.join(a.out, 'MARC_User_Authorization_Path.xlsx'))


def main():
    ap = argparse.ArgumentParser()
    for k in ('date', 'old_dpu', 'old_aup', 'dpu_template', 'aup_template', 'train', 'live', 'watch', 'out'):
        ap.add_argument('--' + k.replace('_', '-'), dest=k, required=True)
    ap.add_argument('--name'); ap.add_argument('--email')
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    tmp = os.path.join(a.out, 'pdf')
    os.makedirs(tmp, exist_ok=True)
    build_dpu(a, os.path.join(tmp, 'MARC_Data_Permission_and_Usage.pdf'))
    build_aup(a, os.path.join(tmp, 'MARC_User_Authorization_Path.pdf'))
    build_xlsx(a, tmp)


if __name__ == '__main__':
    main()
